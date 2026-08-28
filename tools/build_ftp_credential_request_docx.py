from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "docs" / "Bodhi-Mitra_FTP_Credential_Request.docx"

PURPLE = RGBColor(76, 29, 149)
VIOLET = RGBColor(124, 58, 237)
INK = RGBColor(31, 31, 31)
MUTED = RGBColor(92, 83, 116)


def set_run_font(run, size=11, bold=False, italic=False, color=INK):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_page_field(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instr, separate, text, end):
        run._r.append(node)
    set_run_font(run, size=9, color=MUTED)


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.07
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.28)
    set_run_font(p.add_run(text), size=11)
    return p


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.58)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.38)
    section.footer_distance = Inches(0.38)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header_p.paragraph_format.space_after = Pt(0)
    set_run_font(
        header_p.add_run("BODHI-MITRA  |  GAUTAM BUDDHA UNIVERSITY"),
        size=8.5,
        bold=True,
        color=MUTED,
    )

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_before = Pt(0)
    set_run_font(footer_p.add_run("Bodhi-Mitra Access Request  |  Page "), size=8.5, color=MUTED)
    add_page_field(footer_p)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_before = Pt(2)
    kicker.paragraph_format.space_after = Pt(2)
    set_run_font(kicker.add_run("FORMAL APPLICATION"), size=9, bold=True, color=VIOLET)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(3)
    set_run_font(title.add_run("Request for Website FTP Credentials"), size=19, bold=True, color=PURPLE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(8)
    set_run_font(
        subtitle.add_run("Bodhi-Mitra Mental-Health Support Platform"),
        size=11.5,
        italic=True,
        color=MUTED,
    )

    recipient = doc.add_paragraph()
    recipient.paragraph_format.space_after = Pt(4)
    recipient_lines = (
        ("To", True),
        ("The Competent Authority", False),
        ("University Website and IT Administration", False),
        ("Gautam Buddha University", False),
        ("Greater Noida, Uttar Pradesh", False),
    )
    for index, (line, bold) in enumerate(recipient_lines):
        set_run_font(recipient.add_run(line), size=11, bold=bold)
        if index < len(recipient_lines) - 1:
            recipient.add_run().add_break()

    subject = doc.add_paragraph()
    shade_paragraph(subject, "F4F0FF")
    subject.paragraph_format.left_indent = Inches(0.12)
    subject.paragraph_format.right_indent = Inches(0.12)
    subject.paragraph_format.space_before = Pt(5)
    subject.paragraph_format.space_after = Pt(5)
    subject.paragraph_format.line_spacing = 1.05
    set_run_font(subject.add_run("Subject: "), size=11, bold=True, color=PURPLE)
    set_run_font(
        subject.add_run("Request for secure FTP/SFTP credentials for deployment of the Bodhi-Mitra website"),
        size=11,
        bold=True,
    )

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)

    salutation = doc.add_paragraph()
    salutation.paragraph_format.space_after = Pt(4)
    set_run_font(salutation.add_run("Respected Sir/Madam,"), size=11)

    add_body(
        doc,
        "With due respect, we request the competent authority to provide the website access credentials required to deploy and maintain the Bodhi-Mitra mental-health support platform on the hosting facility approved by the University.",
    )
    add_body(
        doc,
        "As the hosting arrangement has already been approved by the authority, the development team now requires controlled access to the designated web directory to upload the production build, publish approved updates, correct deployment issues, and maintain the website without affecting other University services.",
    )
    add_body(
        doc,
        "We kindly request that the following connection information be provided: the server hostname or IP address, approved transfer protocol, port number, project-specific username, initial password or authentication key, assigned website directory or document root, permitted read/write access, domain or subdomain mapping details, and the relevant technical contact for deployment assistance. For security, SFTP or FTPS access is preferred wherever supported, and the password or private access information may kindly be shared through a secure channel rather than ordinary email or an open messaging group.",
    )
    add_body(
        doc,
        "A separate least-privilege account limited to the Bodhi-Mitra directory would be preferable to a shared or administrator-level account. If the University requires IP allow-listing, VPN access, multi-factor authentication, a change-request procedure, scheduled deployment windows, or prior backups before an update, kindly provide the applicable instructions as well.",
    )
    add_body(
        doc,
        "The credentials will be used only by authorised project personnel for official deployment and maintenance work. They will be kept confidential, stored securely, never committed to the source-code repository, and not shared with unauthorised persons. Any password-change, access-log, or credential-rotation policy prescribed by the University will be followed.",
    )
    add_body(
        doc,
        "We therefore request that the required FTP/SFTP credentials and deployment instructions be issued at the earliest convenience so that the approved website deployment can proceed without further delay. We shall be grateful for your consideration and support.",
    )

    closing = doc.add_paragraph()
    closing.paragraph_format.space_before = Pt(3)
    closing.paragraph_format.space_after = Pt(2)
    set_run_font(closing.add_run("Yours sincerely,"), size=11)

    team = doc.add_paragraph()
    team.paragraph_format.space_after = Pt(1)
    set_run_font(team.add_run("Bodhi-Mitra Project Team"), size=11, bold=True, color=PURPLE)

    department = doc.add_paragraph()
    department.paragraph_format.space_after = Pt(5)
    set_run_font(
        department.add_run(
            "Department of Psychology and Mental Health\n"
            "School of Humanities and Social Sciences\n"
            "Gautam Buddha University"
        ),
        size=10.5,
    )

    sign = doc.add_paragraph()
    sign.paragraph_format.space_before = Pt(1)
    sign.paragraph_format.space_after = Pt(0)
    set_run_font(sign.add_run("Date: ____________________"), size=10.5)
    sign.add_run(" " * 14)
    set_run_font(sign.add_run("Authorised Signature: ____________________"), size=10.5)

    doc.core_properties.title = "Bodhi-Mitra FTP Credential Request"
    doc.core_properties.subject = "Request for secure website deployment credentials"
    doc.core_properties.author = "Bodhi-Mitra Project Team"
    doc.core_properties.keywords = "Bodhi-Mitra, FTP, SFTP, credentials, website deployment"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
